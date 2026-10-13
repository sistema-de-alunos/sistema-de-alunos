"""Confere a planilha de medidas: valores em cache, células vazias e conta igual à do app.

Uso: python tools/testar_planilha_medidas.py [pasta_saida]  (precisa de openpyxl).
Se o LibreOffice estiver instalado, também recalcula as fórmulas nele, compara
com o cache e gera um PDF de cada arquivo para conferência visual.
"""
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from core.planilha_medidas import gerar_planilha_medidas, nome_arquivo_medidas  # noqa: E402
from gui.dialogs.cadastro_alunos import _calcular_composicao_jp7  # noqa: E402

CIRCUNFERENCIAS = {
    "ombro": 118.0, "torax": 101.5, "cintura": 84.0, "abdominal": 88.2, "quadril": 99.0,
    "braco_e": 35.0, "braco_e_contraido": 37.5, "antebraco_e": 29.0,
    "braco_d": 35.5, "braco_d_contraido": 38.0, "antebraco_d": 29.4,
    "coxa_e": 58.0, "coxa_d": 58.5, "panturrilha_e": 38.0, "panturrilha_d": 38.2,
}
DOBRAS = {
    "triceps": 12.0, "peito": 10.5, "axilar_media": 11.0, "subescapular": 14.0,
    "dobra_abdominal": 20.0, "supra_iliaca": 15.5, "coxa_dobra": 13.0,
}
CASOS = {
    "masculino_completo": ("Masculino", 30, 80.4, CIRCUNFERENCIAS, DOBRAS),
    "feminino_completo": ("Feminino", 27, 62.0, CIRCUNFERENCIAS, DOBRAS),
    # parcial: falta uma dobra -> composição em branco; algumas circunferências vazias
    "masculino_parcial": ("Masculino", 41, 90.0, {"ombro": 120.0, "cintura": 95.0},
                          {**DOBRAS, "coxa_dobra": None}),
    "feminino_parcial": ("Feminino", 35, None, {"quadril": 101.0}, {}),
}
DATA = date(2026, 10, 13)


def _soffice():
    return shutil.which("soffice") or next(
        (str(p) for p in Path("C:/Program Files").glob("LibreOffice*/program/soffice.exe")), None
    )


def main() -> None:
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp())
    saida.mkdir(parents=True, exist_ok=True)
    assert nome_arquivo_medidas('Ana / Maria: "Silva"?', DATA) == "Medidas_Ana_Maria_Silva_13-10-2026.xlsx"

    for nome, (sexo, idade, peso, circunferencias, dobras) in CASOS.items():
        arquivo = saida / f"{nome}.xlsx"
        gerar_planilha_medidas(arquivo, sexo=sexo, idade=idade, data=DATA,
                               circunferencias=circunferencias, dobras=dobras, peso=peso)

        cache = openpyxl.load_workbook(arquivo, data_only=True).active
        formulas = openpyxl.load_workbook(arquivo).active
        assert all(cache[f"C{linha}"].value.date() == DATA for linha in (2, 21, 33))
        assert cache["C3"].value == circunferencias.get("ombro")
        assert cache["C4"].value == circunferencias.get("torax")
        assert cache["C30"].value == peso
        assert formulas["C38"].value.startswith("=IF(")
        for coluna in "DEF":  # só a data escolhida
            assert all(cache[f"{coluna}{linha}"].value is None for linha in range(2, 39))

        app = _calcular_composicao_jp7(sexo, idade, peso, dobras)
        valores = {linha: cache[f"C{linha}"].value for linha in range(34, 39)}
        if app is None:
            assert all(valor in (None, "") for valor in valores.values()), valores
        else:
            assert abs(valores[38] - app["densidade"]) < 1e-12
            assert abs(valores[37] - app["percentual"]) < 1e-9
            assert abs(valores[34] - app["massa_magra"]) < 1e-9
            assert abs(valores[35] - app["massa_gorda"]) < 1e-9
            assert abs(valores[36] - app["percentual_protegido"]) < 1e-9
        print("ok:", arquivo, {linha: valor for linha, valor in valores.items()})

    soffice = _soffice()
    if soffice is None:
        print("LibreOffice não encontrado: pulando recálculo e PDF.")
        return
    recalculado = saida / "recalculado"
    for formato, pasta in (("xlsx", recalculado), ("pdf", saida)):
        subprocess.run([soffice, "--headless", "--convert-to", formato, "--outdir", str(pasta),
                        *[str(saida / f"{nome}.xlsx") for nome in CASOS]], check=True,
                       capture_output=True)
    for nome in CASOS:
        cache = openpyxl.load_workbook(saida / f"{nome}.xlsx", data_only=True).active
        recalc = openpyxl.load_workbook(recalculado / f"{nome}.xlsx", data_only=True).active
        for linha in range(34, 39):
            esperado, obtido = cache[f"C{linha}"].value, recalc[f"C{linha}"].value
            if isinstance(esperado, float):
                assert abs(esperado - obtido) < 1e-9, (nome, linha, esperado, obtido)
            else:
                assert obtido in (None, ""), (nome, linha, obtido)
    print("LibreOffice recalculou igual ao cache. PDFs em", saida)


if __name__ == "__main__":
    main()
