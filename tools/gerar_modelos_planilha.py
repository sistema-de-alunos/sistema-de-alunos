"""Gera os modelos compactos de planilha de medidas a partir dos modelos originais.

Uso: python tools/gerar_modelos_planilha.py  (precisa de openpyxl e Pillow, ver
tools/requirements.txt). Rode de novo só se os modelos em
assets/planilhas/originais mudarem.

Mantém apenas as seções Circunferência, Dobras Cutâneas e Composição Corporal
(linhas 55-92 do original) com o mesmo visual, sem o resto do modelo. Os
valores e fórmulas da coluna C são preenchidos pelo app (core/planilha_medidas.py).
"""
import io
import re
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as ImagemPlanilha
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image

RAIZ = Path(__file__).resolve().parent.parent
PASTA = RAIZ / "assets" / "planilhas"
# Mesmo boneco da tabela de Circunferências do app, no lugar do boneco do modelo.
CORPO_DO_APP = {
    "MASCULINO": RAIZ / "assets" / "corpos" / "corpohomem.png",
    "FEMININO": RAIZ / "assets" / "corpos" / "corpomulher.png",
}
AREA_CORPO = ("G", "I", 2, 17)  # células mescladas G2:I17
MARGEM_CORPO_PX = 6

PRIMEIRA_LINHA = 55  # "Circunferência" no original
ULTIMA_LINHA = 92  # "Densidade Corporal" no original
DESLOCAMENTO = PRIMEIRA_LINHA - 1
LINHAS_DATAS = (2, 21, 33)  # já nas posições novas
LINHAS_FORMULAS = range(34, 39)
LARGURA_COLUNA_DATA = 12


def _imagens_das_secoes(caminho: Path):
    """Imagens ancoradas dentro das seções mantidas, já recortadas como no Excel."""
    with zipfile.ZipFile(caminho) as arquivo:
        desenho = arquivo.read("xl/drawings/drawing1.xml").decode("utf-8")
        relacoes = arquivo.read("xl/drawings/_rels/drawing1.xml.rels").decode("utf-8")
        for ancora in re.findall(r"<xdr:twoCellAnchor.*?</xdr:twoCellAnchor>", desenho, re.S):
            if "<xdr:pic>" not in ancora:
                continue
            de, ate = (
                [int(n) for n in re.search(
                    rf"<xdr:{ponta}><xdr:col>(\d+)</xdr:col><xdr:colOff>(\d+)</xdr:colOff>"
                    rf"<xdr:row>(\d+)</xdr:row><xdr:rowOff>(\d+)</xdr:rowOff>", ancora
                ).groups()]
                for ponta in ("from", "to")
            )
            if not PRIMEIRA_LINHA - 1 <= de[2] <= ULTIMA_LINHA - 1:
                continue
            rid = re.search(r'r:embed="(\w+)"', ancora).group(1)
            alvo = re.search(rf'<Relationship[^>]*Id="{rid}"[^>]*>', relacoes).group(0)
            midia = re.search(r'Target="\.\./([^"]+)"', alvo).group(1)
            imagem = Image.open(io.BytesIO(arquivo.read(f"xl/{midia}")))

            recorte = dict(re.findall(r'(\w)="(\d+)"', (re.findall(r"<a:srcRect[^>]*>", ancora) or [""])[0]))
            largura, altura = imagem.size
            imagem = imagem.crop((
                largura * int(recorte.get("l", 0)) // 100000,
                altura * int(recorte.get("t", 0)) // 100000,
                largura - largura * int(recorte.get("r", 0)) // 100000,
                altura - altura * int(recorte.get("b", 0)) // 100000,
            ))
            buffer = io.BytesIO()
            imagem.save(buffer, format="PNG")
            yield buffer, de, ate


def _largura_coluna_px(planilha, letra: str) -> int:
    return int(planilha.column_dimensions[letra].width * 7 + 5)


def _altura_linha_px(planilha, linha: int) -> int:
    return round(planilha.row_dimensions[linha].height * 96 / 72)


def _adicionar_corpo_do_app(planilha, caminho: Path) -> None:
    """Centraliza o boneco do app na área do boneco, mantendo a proporção."""
    coluna_inicial, coluna_final, linha_inicial, linha_final = AREA_CORPO
    colunas = [chr(c) for c in range(ord(coluna_inicial), ord(coluna_final) + 1)]
    linhas = range(linha_inicial, linha_final + 1)
    largura_area = sum(_largura_coluna_px(planilha, letra) for letra in colunas)
    altura_area = sum(_altura_linha_px(planilha, linha) for linha in linhas)

    imagem = Image.open(caminho)
    imagem = imagem.crop(imagem.getbbox())
    escala = min(
        (largura_area - 2 * MARGEM_CORPO_PX) / imagem.width,
        (altura_area - 2 * MARGEM_CORPO_PX) / imagem.height,
    )
    largura, altura = round(imagem.width * escala), round(imagem.height * escala)
    imagem.thumbnail((largura * 3, altura * 3))  # 3x para ficar nítido ao dar zoom/imprimir
    buffer = io.BytesIO()
    imagem.save(buffer, format="PNG", optimize=True)

    # Âncora: célula onde cai o canto superior esquerdo + deslocamento dentro dela.
    x, y = (largura_area - largura) // 2, (altura_area - altura) // 2
    indice_coluna = ord(coluna_inicial) - ord("A")
    for letra in colunas:
        if x < _largura_coluna_px(planilha, letra):
            break
        x -= _largura_coluna_px(planilha, letra)
        indice_coluna += 1
    indice_linha = linha_inicial - 1
    for linha in linhas:
        if y < _altura_linha_px(planilha, linha):
            break
        y -= _altura_linha_px(planilha, linha)
        indice_linha += 1

    figura = ImagemPlanilha(buffer)
    figura.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=indice_coluna, colOff=pixels_to_EMU(x), row=indice_linha, rowOff=pixels_to_EMU(y)
        ),
        ext=XDRPositiveSize2D(pixels_to_EMU(largura), pixels_to_EMU(altura)),
    )
    planilha.add_image(figura)


def gerar(original: Path, destino: Path, corpo_do_app: Path) -> None:
    livro = openpyxl.load_workbook(original)
    planilha = livro["Planilha1"]
    planilha._images = []
    planilha._charts = []

    mesclagens = [
        faixa for faixa in planilha.merged_cells.ranges
        if faixa.min_row >= PRIMEIRA_LINHA and faixa.max_row <= ULTIMA_LINHA and faixa.max_col <= 9
    ]
    alturas = {
        linha: (planilha.row_dimensions[linha].height, planilha.row_dimensions[linha].hidden)
        for linha in range(PRIMEIRA_LINHA, ULTIMA_LINHA + 1)
    }
    for faixa in list(planilha.merged_cells.ranges):
        planilha.merged_cells.remove(faixa)

    planilha.delete_rows(ULTIMA_LINHA + 1, planilha.max_row)
    planilha.delete_rows(1, DESLOCAMENTO)
    planilha.delete_cols(10, planilha.max_column)

    for faixa in mesclagens:
        faixa.shift(row_shift=-DESLOCAMENTO)
        planilha.merge_cells(faixa.coord)
    planilha.row_dimensions.clear()
    for linha, (altura, oculta) in alturas.items():
        planilha.row_dimensions[linha - DESLOCAMENTO].height = altura
        planilha.row_dimensions[linha - DESLOCAMENTO].hidden = oculta

    for linha in LINHAS_DATAS:
        for coluna in "CDEF":
            planilha[f"{coluna}{linha}"].number_format = "dd/mm/yyyy"
    # O modelo usa "d-mmm"; dd/mm/aaaa não cabe na largura original. No modelo a
    # largura de C vale para o grupo C:F inteiro.
    colunas_datas = planilha.column_dimensions["C"]
    assert (colunas_datas.min, colunas_datas.max) == (3, 6), (colunas_datas.min, colunas_datas.max)
    colunas_datas.width = LARGURA_COLUNA_DATA
    planilha.oddHeader.center.text = None  # cabeçalho do original é só uma quebra de linha
    for linha in LINHAS_FORMULAS:
        for coluna in "CDEF":
            planilha[f"{coluna}{linha}"].value = None

    _adicionar_corpo_do_app(planilha, corpo_do_app)
    for buffer, de, ate in _imagens_das_secoes(original):
        if de[2] - DESLOCAMENTO < AREA_CORPO[3]:  # boneco do modelo: trocado pelo do app
            continue
        imagem = ImagemPlanilha(buffer)
        imagem.anchor = TwoCellAnchor(
            editAs="oneCell",
            _from=AnchorMarker(col=de[0], colOff=de[1], row=de[2] - DESLOCAMENTO, rowOff=de[3]),
            to=AnchorMarker(col=ate[0], colOff=ate[1], row=ate[2] - DESLOCAMENTO, rowOff=ate[3]),
        )
        planilha.add_image(imagem)

    planilha.sheet_view.selection[0].activeCell = "A1"
    planilha.sheet_view.selection[0].sqref = "A1"
    planilha.print_area = f"A1:I{ULTIMA_LINHA - DESLOCAMENTO}"
    livro.calculation.fullCalcOnLoad = True
    livro.save(destino)


def main() -> None:
    for sexo in ("MASCULINO", "FEMININO"):
        destino = PASTA / f"medidas_{sexo.lower()}.xlsx"
        gerar(PASTA / "originais" / f"ADIPOMETRO-{sexo}.xlsx", destino, CORPO_DO_APP[sexo])
        print("gerado:", destino)


if __name__ == "__main__":
    main()
